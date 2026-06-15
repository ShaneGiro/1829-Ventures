"""Dealroom CSV/Excel parsing and normalization.

Dealroom exports include metadata rows before the real header, then store many
one-to-many values as semicolon-delimited arrays inside individual cells. Exports
arrive as CSV or modern Excel (.xlsx/.xlsm); both are normalized into the same
``DealroomParseResult`` here.

Column names are not hard-coded in this module — they live in
``app.integrations.dealroom_columns`` (the single source of truth), so a Dealroom
header rename is a one-line edit there, not a change to parsing logic.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any
from urllib.parse import urlparse

from app.core.constants import DEFAULT_SECTOR
from app.integrations.dealroom_columns import (
    HEADER_MARKERS,
    SUPPORTED_UPLOAD_EXTENSIONS,
    TEMPLATE_COLUMNS,
    DealroomColumn,
)

MISSING_VALUES = {"", "n/a", "na", "null", "none", "-"}


@dataclass(frozen=True)
class DealroomFounder:
    name: str
    status: str | None = None
    gender: str | None = None
    university: str | None = None
    linkedin_url: str | None = None
    background: str | None = None


@dataclass(frozen=True)
class DealroomFundingRound:
    round_type: str | None
    amount: Decimal | None
    currency: str | None
    date: str | None
    investors: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class DealroomParsedRow:
    row_number: int
    raw: dict[str, str]
    dealroom_id: str | None
    name: str | None
    dealroom_url: str | None
    website: str | None
    domain: str | None
    tagline: str | None
    description: str | None
    stage: str | None
    sector: str
    sector_warnings: list[str]
    industries: list[str]
    sub_industries: list[str]
    tags: list[str]
    city: str | None
    state: str | None
    country: str | None
    latitude: Decimal | None
    longitude: Decimal | None
    founders: list[DealroomFounder]
    funding_rounds: list[DealroomFundingRound]


@dataclass(frozen=True)
class DealroomParseResult:
    header_row_number: int
    metadata_rows: list[list[str]]
    headers: list[str]
    rows: list[DealroomParsedRow]


class UnsupportedDealroomFile(ValueError):
    """Raised when an uploaded file is not a supported Dealroom export format."""


def parse_dealroom_file(content: bytes, filename: str | None) -> DealroomParseResult:
    """Parse a Dealroom export (CSV or Excel) into normalized rows.

    Dispatches on the file extension. Both formats are reduced to a grid of string
    cells (``list[list[str]]``) and then run through the same parsing core, so the
    metadata-row skipping, header detection, and field mapping behave identically.
    """
    extension = _extension(filename)
    if extension in {".xlsx", ".xlsm"}:
        raw_rows = _read_excel_rows(content)
    elif extension == ".csv":
        raw_rows = _read_csv_rows(content)
    else:
        supported = ", ".join(SUPPORTED_UPLOAD_EXTENSIONS)
        raise UnsupportedDealroomFile(
            f"Unsupported file type '{extension or filename}'. Supported formats: {supported}."
        )
    return _build_result(raw_rows)


def parse_dealroom_csv(content: bytes | str) -> DealroomParseResult:
    """Parse a Dealroom CSV export into normalized rows."""
    return _build_result(_read_csv_rows(content))


def build_template_csv() -> str:
    """Render the canonical Dealroom column template as CSV text.

    Produces a header-only CSV listing every column from the most recent export
    template, so the team can see/share the exact expected column names. The
    output re-imports cleanly (header detection finds the row; there are no data
    rows). Columns come from the registry, so the template tracks schema edits.
    """
    buffer = io.StringIO()
    csv.writer(buffer).writerow(TEMPLATE_COLUMNS)
    return buffer.getvalue()


def _build_result(raw_rows: list[list[str]]) -> DealroomParseResult:
    header_index = _find_header_index(raw_rows)
    headers = [header.strip() for header in raw_rows[header_index]]

    parsed_rows: list[DealroomParsedRow] = []
    for index, row in enumerate(raw_rows[header_index + 1 :], start=header_index + 2):
        raw = _row_to_dict(headers, row)
        if not any(value.strip() for value in raw.values()):
            continue
        parsed_rows.append(_parse_row(index, raw))

    return DealroomParseResult(
        header_row_number=header_index + 1,
        metadata_rows=raw_rows[:header_index],
        headers=headers,
        rows=parsed_rows,
    )


def _extension(filename: str | None) -> str:
    if not filename:
        return ""
    return PurePosixPath(filename).suffix.lower()


def _read_csv_rows(content: bytes | str) -> list[list[str]]:
    text = content.decode("utf-8-sig") if isinstance(content, bytes) else content
    return list(csv.reader(io.StringIO(text)))


def _read_excel_rows(content: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnsupportedDealroomFile(
            "Excel support requires the 'openpyxl' package to be installed."
        ) from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            return []
        return [
            [_cell_to_str(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _cell_to_str(value: Any) -> str:
    """Render an Excel cell as the string a CSV cell would have held."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        # Excel stores all numbers as floats; render integer-valued ones without a
        # trailing ".0" so IDs and years survive (e.g. 4936908.0 -> "4936908").
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def split_semicolon(value: str | None) -> list[str | None]:
    """Split Dealroom's semicolon arrays while preserving empty positions."""
    if value is None:
        return []
    if value == "":
        return []
    return [_clean_value(part) for part in value.split(";")]


def normalize_domain(website: str | None) -> str | None:
    value = _clean_value(website)
    if value is None:
        return None
    parsed = urlparse(value if "://" in value else f"https://{value}")
    host = (parsed.netloc or parsed.path).lower()
    host = host.removeprefix("www.")
    return host.rstrip("/") or None


def map_dealroom_sector(industries: list[str], sub_industries: list[str]) -> tuple[str, list[str]]:
    """Map Dealroom industries to 1829's five-sector taxonomy."""
    haystack = " ".join([*industries, *sub_industries]).lower()
    rules: tuple[tuple[str, tuple[str, ...]], ...] = (
        (
            "Photonics, Imaging & Quantum",
            ("photon", "imaging", "quantum", "optics", "sensor", "semiconductor"),
        ),
        (
            "Clean Tech & Energy",
            (
                "clean",
                "climate",
                "energy",
                "battery",
                "carbon",
                "solar",
                "recycling",
                "sustainab",
                "green",
            ),
        ),
        (
            "Life Sciences & Health Tech",
            (
                "health",
                "medical",
                "bio",
                "life science",
                "pharma",
                "therapeutic",
                "diagnostic",
                "patient",
            ),
        ),
        (
            "Intelligent Systems, AI & Cyber",
            (
                "artificial intelligence",
                "machine learning",
                "robot",
                "cyber",
                "security",
                "software",
                "enterprise",
                "automation",
                "data",
                "defence",
                "defense",
                "drone",
            ),
        ),
    )
    for sector, needles in rules:
        if any(needle in haystack for needle in needles):
            return sector, []
    if industries or sub_industries:
        return DEFAULT_SECTOR, ["unmapped_dealroom_taxonomy"]
    return DEFAULT_SECTOR, ["missing_dealroom_taxonomy"]


def _find_header_index(rows: list[list[str]]) -> int:
    for index, row in enumerate(rows):
        normalized = {cell.strip() for cell in row}
        if all(marker in normalized for marker in HEADER_MARKERS):
            return index
    raise ValueError("Could not find Dealroom header row")


def _row_to_dict(headers: list[str], row: list[str]) -> dict[str, str]:
    values = [*row, *([""] * max(len(headers) - len(row), 0))]
    return dict(zip(headers, values[: len(headers)], strict=False))


def _parse_row(row_number: int, raw: dict[str, str]) -> DealroomParsedRow:
    industries = _clean_list(split_semicolon(raw.get(DealroomColumn.INDUSTRIES)))
    sub_industries = _clean_list(split_semicolon(raw.get(DealroomColumn.SUB_INDUSTRIES)))
    sector, sector_warnings = map_dealroom_sector(industries, sub_industries)
    website = _clean_value(raw.get(DealroomColumn.WEBSITE))
    tagline = _clean_value(raw.get(DealroomColumn.TAGLINE))
    return DealroomParsedRow(
        row_number=row_number,
        raw=raw,
        dealroom_id=_clean_value(raw.get(DealroomColumn.ID)),
        name=_clean_value(raw.get(DealroomColumn.NAME)),
        dealroom_url=_clean_value(raw.get(DealroomColumn.DEALROOM_URL)),
        website=website,
        domain=normalize_domain(website),
        tagline=tagline,
        description=_clean_value(raw.get(DealroomColumn.LONG_DESCRIPTION)) or tagline,
        stage=_clean_value(raw.get(DealroomColumn.GROWTH_STAGE))
        or _clean_value(raw.get(DealroomColumn.LAST_ROUND)),
        sector=sector,
        sector_warnings=sector_warnings,
        industries=industries,
        sub_industries=sub_industries,
        tags=_clean_list(
            split_semicolon(raw.get(DealroomColumn.ALL_TAGS) or raw.get(DealroomColumn.TAGS))
        ),
        city=_clean_value(raw.get(DealroomColumn.HQ_CITY)),
        state=_clean_value(raw.get(DealroomColumn.HQ_STATE)),
        country=_clean_value(raw.get(DealroomColumn.HQ_COUNTRY)),
        latitude=_decimal_or_none(raw.get(DealroomColumn.LATITUDE)),
        longitude=_decimal_or_none(raw.get(DealroomColumn.LONGITUDE)),
        founders=_parse_founders(raw),
        funding_rounds=_parse_funding_rounds(raw),
    )


def _parse_founders(raw: dict[str, str]) -> list[DealroomFounder]:
    names = split_semicolon(raw.get(DealroomColumn.FOUNDERS))
    statuses = split_semicolon(raw.get(DealroomColumn.FOUNDERS_STATUSES))
    genders = split_semicolon(raw.get(DealroomColumn.FOUNDERS_GENDERS))
    universities = split_semicolon(raw.get(DealroomColumn.FOUNDERS_UNIVERSITIES))
    linkedin_urls = split_semicolon(raw.get(DealroomColumn.FOUNDERS_LINKEDIN))
    backgrounds = split_semicolon(raw.get(DealroomColumn.FOUNDERS_BACKGROUNDS))

    founders: list[DealroomFounder] = []
    for index, name in enumerate(names):
        if name is None:
            continue
        founders.append(
            DealroomFounder(
                name=name,
                status=_at(statuses, index),
                gender=_at(genders, index),
                university=_at(universities, index),
                linkedin_url=_at(linkedin_urls, index),
                background=_at(backgrounds, index),
            )
        )
    return founders


def _parse_funding_rounds(raw: dict[str, str]) -> list[DealroomFundingRound]:
    round_types = split_semicolon(raw.get(DealroomColumn.EACH_ROUND_TYPE))
    amounts = split_semicolon(raw.get(DealroomColumn.EACH_ROUND_AMOUNT))
    currencies = split_semicolon(raw.get(DealroomColumn.EACH_ROUND_CURRENCY))
    dates = split_semicolon(raw.get(DealroomColumn.EACH_ROUND_DATE))
    investors = split_semicolon(raw.get(DealroomColumn.EACH_ROUND_INVESTORS))
    count = max(len(round_types), len(amounts), len(currencies), len(dates), len(investors))

    rounds: list[DealroomFundingRound] = []
    for index in range(count):
        round_type = _at(round_types, index)
        amount = _decimal_or_none(_at(amounts, index))
        currency = _at(currencies, index)
        date = _at(dates, index)
        investor_list = _clean_list((_at(investors, index) or "").split("++"))
        if round_type or amount is not None or currency or date or investor_list:
            rounds.append(
                DealroomFundingRound(
                    round_type=round_type,
                    amount=amount,
                    currency=currency,
                    date=date,
                    investors=investor_list,
                )
            )
    return rounds


def _at(values: list[str | None], index: int) -> str | None:
    return values[index] if index < len(values) else None


def _clean_list(values: Iterable[str | None]) -> list[str]:
    return [value for value in (_clean_value(value) for value in values) if value is not None]


def _clean_value(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = re.sub(r"\s+", " ", value).strip()
    return None if cleaned.lower() in MISSING_VALUES else cleaned


def _decimal_or_none(value: str | None) -> Decimal | None:
    cleaned = _clean_value(value)
    if cleaned is None:
        return None
    try:
        return Decimal(cleaned.replace(",", ""))
    except InvalidOperation:
        return None
